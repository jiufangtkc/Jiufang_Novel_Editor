import os
import sys
import tempfile
import unittest
import json
from PyQt6.QtWidgets import QApplication
from PyQt6.QtCore import Qt
from docx import Document

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))

from utils.dataset_formatter import DatasetFormatter
from utils.markdown_converter import MarkdownConverter
from views.dialogs.dataset_export_dialog import DatasetExportDialog
from views.main_window import MainWindow
from controllers.main_controller import MainController

app = QApplication.instance()
if not app:
    app = QApplication(sys.argv)


class TestDatasetExport(unittest.TestCase):
    """測試設定資料集匯出精靈與格式化轉換功能。"""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_dataset_formatter_depth_degradation(self):
        """驗證 DatasetFormatter 在深度超過 3 層時正確觸發標題降級與麵包屑路徑。"""
        # 建立深度高達 5 層的卡片結構
        cards_data = {
            "character": [
                {
                    "title": "主角陣營",
                    "id": "c1",
                    "content": "故事核心陣營",
                    "children": [
                        {
                            "title": "勇者亞倫",
                            "id": "c2",
                            "content": "手持聖劍的傳奇少年",
                            "children": [
                                {
                                    "title": "身世秘密",
                                    "id": "c3",
                                    "content": "實為古代王國失落的王儲",
                                    "children": [
                                        {
                                            "title": "龍之胎記",
                                            "id": "c4",
                                            "content": "右臂上浮現古代龍紋，危機時會發光",
                                            "children": []
                                        }
                                    ]
                                }
                            ]
                        }
                    ]
                }
            ]
        }

        md_output = DatasetFormatter.format_to_markdown(
            cards_data=cards_data,
            categories_meta={"character": "登場角色"},
            book_title="星空之誓"
        )

        # 1. 驗證總標題與大分類 (Level 1)
        self.assertIn("# 星空之誓 — 設定資料集", md_output)
        self.assertIn("# 登場角色", md_output)

        # 2. 驗證正常標題層級 (Level 2 & Level 3)
        self.assertIn("## 主角陣營", md_output)
        self.assertIn("### 勇者亞倫", md_output)

        # 3. 核心檢驗：絕對不得產出 Level 4 (####) 或 Level 5 (#####)
        self.assertNotIn("#### 身世秘密", md_output)
        self.assertNotIn("##### 龍之胎記", md_output)

        # 4. 驗證降級排版：使用麵包屑清單與引言
        self.assertIn("* **[主角陣營 > 勇者亞倫] 身世秘密**", md_output)
        self.assertIn("> 實為古代王國失落的王儲", md_output)
        self.assertIn("* **[主角陣營 > 勇者亞倫 > 身世秘密] 龍之胎記**", md_output)
        self.assertIn("> 右臂上浮現古代龍紋，危機時會發光", md_output)

    def test_dataset_export_dialog_filtering(self):
        """驗證 DatasetExportDialog 能夠正確依據勾選過濾卡片資料。"""
        cards_data = {
            "character": [
                {
                    "title": "主角",
                    "id": "c_hero",
                    "content": "第一主角",
                    "children": []
                },
                {
                    "title": "配角",
                    "id": "c_sub",
                    "content": "同伴",
                    "children": []
                }
            ],
            "world": [
                {
                    "title": "魔法體系",
                    "id": "w_magic",
                    "content": "五大元素",
                    "children": []
                }
            ]
        }

        dialog = DatasetExportDialog(cards_data=cards_data)
        
        # 初始狀態：全選
        filtered_all = dialog.get_filtered_cards_data()
        self.assertIn("character", filtered_all)
        self.assertIn("world", filtered_all)
        self.assertEqual(len(filtered_all["character"]), 2)

        # 取消勾選世界觀分類
        for i in range(dialog.tree_widget.topLevelItemCount()):
            item = dialog.tree_widget.topLevelItem(i)
            if item.text(0) == "世界觀":
                item.setCheckState(0, Qt.CheckState.Unchecked)
                dialog._on_item_changed(item, 0)

        filtered_partial = dialog.get_filtered_cards_data()
        self.assertIn("character", filtered_partial)
        self.assertNotIn("world", filtered_partial)

        dialog.close()

    def test_dataset_render_to_docx(self):
        """驗證 DatasetFormatter 產出的 Markdown 可成功轉為 DOCX 文件。"""
        cards_data = {
            "world": [
                {
                    "title": "大陸地理",
                    "id": "w1",
                    "content": "位於艾爾斯大陸中部，四季分明。",
                    "children": []
                }
            ]
        }
        md_text = DatasetFormatter.format_to_markdown(cards_data=cards_data)
        doc = Document()
        MarkdownConverter.render_to_docx(md_text, doc)

        docx_path = os.path.join(self.temp_dir.name, "dataset.docx")
        doc.save(docx_path)
        self.assertTrue(os.path.exists(docx_path))

        # 重新讀取驗證段落存在
        reloaded = Document(docx_path)
        texts = [p.text for p in reloaded.paragraphs]
        self.assertTrue(any("世界觀" in t for t in texts))
        self.assertTrue(any("大陸地理" in t for t in texts))
        self.assertTrue(any("艾爾斯大陸中部" in t for t in texts))

    def test_menu_actions_and_shortcuts(self):
        """驗證選單上的匯出選項更名與新功能按鈕。"""
        view = MainWindow()
        mc = MainController(view)

        # 驗證原本的匯出已更名
        self.assertEqual(view.action_export.text(), "匯出小說文本(&E)...")
        self.assertEqual(view.action_export.shortcut().toString(), "Ctrl+E")

        # 驗證新增的匯出設定資料集動作
        self.assertTrue(hasattr(view, "action_export_dataset"))
        self.assertEqual(view.action_export_dataset.text(), "匯出設定資料集(&D)...")
        self.assertEqual(view.action_export_dataset.shortcut().toString(), "Ctrl+Shift+E")

        # 清理計時器
        if hasattr(mc, 'writing_timer') and mc.writing_timer.isActive():
            mc.writing_timer.stop()
        if hasattr(mc, 'auto_save_timer') and mc.auto_save_timer.isActive():
            mc.auto_save_timer.stop()
        view.close()
